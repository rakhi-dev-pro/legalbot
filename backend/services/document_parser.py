import io
import re
import logging
from typing import List, Dict, Any, Tuple
import fitz  # PyMuPDF
import docx

try:
    import pytesseract
    from PIL import Image
except ImportError:
    pytesseract = None
    Image = None

logger = logging.getLogger("legalbot.document_parser")

# A page counts as "text-native" once its extractable text clears this
# density (characters per square inch of page area). Using density instead
# of a raw character count means the threshold scales correctly across
# page sizes (A4, US Letter, Legal, oddly-cropped scans).
MIN_TEXT_DENSITY_CHARS_PER_SQIN = 6.0

# Fraction of page area that embedded images must cover before a page is
# treated as "image-dominant" (a scan, photo, or a flattened signature/stamp/seal block)
IMAGE_DOMINANCE_AREA_RATIO = 0.35


def _page_area_sq_inches(page: "fitz.Page") -> float:
    rect = page.rect
    # PDF units are points (1/72 inch)
    return max((rect.width / 72.0) * (rect.height / 72.0), 0.01)


def _image_coverage_ratio(page: "fitz.Page") -> float:
    """
    Fraction of the page covered by embedded raster images, based on each
    image's actual placement rectangle on the page rather than raw pixel dimensions.
    """
    page_area = _page_area_sq_inches(page)
    covered = 0.0
    try:
        for img in page.get_images(full=True):
            xref = img[0]
            for rect in page.get_image_rects(xref):
                covered += (rect.width / 72.0) * (rect.height / 72.0)
    except Exception:
        return 0.0
    return min(covered / page_area, 1.0)


def ocr_page_pixmap(page: "fitz.Page", dpi: int = 200) -> str:
    """
    Perform full high-resolution Tesseract OCR on a PDF page pixmap image.
    Used when a page's embedded text layer is missing or sparse.
    """
    if not pytesseract or not Image:
        logger.warning("pytesseract or PIL is not available in environment for OCR fallback.")
        return ""

    try:
        pix = page.get_pixmap(dpi=dpi)
        pil_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        ocr_text = pytesseract.image_to_string(pil_img).strip()
        return ocr_text
    except Exception as e:
        logger.warning(f"Failed to perform Tesseract OCR on page pixmap: {e}")
        return ""


def clean_extracted_text(text: str) -> str:
    """
    Clean OCR/extracted document text by stripping out eSign watermark noise,
    provider hash tokens, notary seal artifacts, and fixing common OCR typos.
    """
    if not text:
        return ""

    lines = text.splitlines()
    cleaned_lines = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            cleaned_lines.append("")
            continue

        # 1. Skip standalone eSign tracking hashes (e.g. 01O98327498273498237)
        if len(stripped.split()) <= 2 and len(stripped) >= 20 and re.search(r'\d', stripped) and re.search(r'[A-Za-z]', stripped):
            continue

        if re.search(r'Szz0H|rreDF|Orie\s+rre', stripped):
            continue
        if re.search(r'(?:eSigned|Signed)\s+using\s+Aadhaar', stripped, re.IGNORECASE):
            continue
        if re.search(r'^Date:\s+[A-Za-z]{3}\s+[A-Za-z]{3}\s+\d+', stripped):
            continue

        # 2. Skip OCR top/bottom margin noise and scanner artifact lines
        if re.search(r'(?:oe\s+)?br\s+Pe\s+ae\s+cecitet', stripped, re.IGNORECASE):
            continue
        if re.search(r'^[a-z]\s+\d+(?:\s+his\s*/)?$', stripped, re.IGNORECASE):
            continue
        if re.match(r'^(?:oe|ee|e)\b$', stripped, re.IGNORECASE):
            continue

        # 3. Fix common OCR typos and remove bullet noise in legal terms
        line_clean = re.sub(r'%(\d+)', r'₹\1', stripped)
        line_clean = re.sub(r'\b10"\s+day\b', '10th day', line_clean)
        line_clean = re.sub(r'\bérsons\b', 'persons', line_clean)
        line_clean = re.sub(r'7<«\s*', '', line_clean)
        line_clean = re.sub(r'^\s*ee\s+', '', line_clean)
        line_clean = re.sub(r'^\s*e\s+(?=Mrs\.|Mr\.|THIS|NOW)', '', line_clean)

        if line_clean.strip():
            cleaned_lines.append(line_clean.strip())

    # Reconstruct text with clean line breaks & normalize multi-newlines
    cleaned_text = "\n".join(cleaned_lines)
    cleaned_text = re.sub(r'\n{3,}', '\n\n', cleaned_text)
    return cleaned_text.strip()


def _needs_ocr(page: "fitz.Page", direct_text: str) -> bool:
    """
    Decide whether a page needs OCR using structural signals -- text
    density and image coverage -- that hold across document types and vendors.
    """
    area = _page_area_sq_inches(page)
    density = len(direct_text) / area if area else 0.0

    if density < MIN_TEXT_DENSITY_CHARS_PER_SQIN:
        return True

    if _image_coverage_ratio(page) >= IMAGE_DOMINANCE_AREA_RATIO:
        return True

    return False


def parse_pdf_bytes(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Extract plain text and page-by-page text map from raw PDF file bytes using PyMuPDF.
    Automatically detects scanned/image pages or sparse eSigned pages using geometry
    and density metrics, and merges Tesseract OCR results with direct text.
    """
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages_list = []
    full_text_chunks = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        direct_text = page.get_text("text").strip()
        cleaned_direct = clean_extracted_text(direct_text)
        final_page_text = cleaned_direct

        if _needs_ocr(page, cleaned_direct):
            logger.info(f"Page {page_num + 1}: Triggering geometry-driven OCR fallback (direct_chars={len(cleaned_direct)})...")
            ocr_text = ocr_page_pixmap(page, dpi=200)
            cleaned_ocr = clean_extracted_text(ocr_text)

            if cleaned_ocr:
                if not cleaned_direct:
                    final_page_text = cleaned_ocr
                elif cleaned_ocr not in cleaned_direct:
                    # Merge native text layer with OCR text
                    final_page_text = f"{cleaned_direct}\n\n{cleaned_ocr}"
                logger.info(f"Page {page_num + 1}: OCR extracted {len(cleaned_ocr)} characters.")

        final_page_text = clean_extracted_text(final_page_text)

        if final_page_text:
            pages_list.append({
                "page_number": page_num + 1,
                "text": final_page_text
            })
            full_text_chunks.append(final_page_text)

    doc.close()
    full_text = "\n\n".join(full_text_chunks)
    return full_text, pages_list


def parse_docx_bytes(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Extract plain text from Microsoft Word (.docx) file bytes using python-docx.
    Extracts document paragraphs and table cells.
    """
    doc_file = io.BytesIO(file_bytes)
    document = docx.Document(doc_file)

    extracted_lines = []

    # 1. Extract Paragraphs
    for p in document.paragraphs:
        txt = p.text.strip()
        if txt:
            extracted_lines.append(txt)

    # 2. Extract Tables
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                extracted_lines.append(row_text)

    full_text = "\n\n".join(extracted_lines)
    pages_list = [{"page_number": 1, "text": full_text}]
    return full_text, pages_list


def extract_document_text(file_bytes: bytes, file_type: str) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Unified entrypoint for document text parsing.
    Supports PDF (with geometry-driven Tesseract OCR for scanned/eSigned images) and DOCX.
    """
    file_type_upper = file_type.upper()
    if file_type_upper == "PDF":
        return parse_pdf_bytes(file_bytes)
    elif file_type_upper in ("DOCX", "DOC"):
        return parse_docx_bytes(file_bytes)
    else:
        raise ValueError(f"Unsupported document format: {file_type}. Only PDF and DOCX are supported.")
